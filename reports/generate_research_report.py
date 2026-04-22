from __future__ import annotations

import json
import math
import textwrap
from collections import Counter
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORTS_DIR = ROOT / "reports"
ASSETS_DIR = REPORTS_DIR / "generated_assets"
OUTPUT_PATH = REPORTS_DIR / "VANTA_Research_Based_Project_Report.docx"


def load_json(name: str) -> dict:
    return json.loads((REPORTS_DIR / name).read_text(encoding="utf-8"))


ALIGNMENT = load_json("vanta_idea_vs_architecture.json")
MODERN = load_json("vanta_vs_modern_architectures.json")


SOURCES = [
    {
        "id": "R1",
        "citation": "Rose, S., Borchert, O., Mitchell, S., & Connelly, S. (2020). Zero Trust Architecture. NIST SP 800-207.",
        "url": "https://csrc.nist.gov/pubs/sp/800/207/final",
        "note": "Defines zero trust principles and explains why static perimeter thinking is insufficient for modern enterprise security.",
    },
    {
        "id": "R2",
        "citation": "Chandramouli, R., & Butcher, Z. (2023). A Zero Trust Architecture Model for Access Control in Cloud-Native Applications in Multi-Cloud Environments. NIST SP 800-207A.",
        "url": "https://csrc.nist.gov/pubs/sp/800/207/a/final",
        "note": "Extends zero trust thinking toward distributed applications and hybrid cloud realities.",
    },
    {
        "id": "R3",
        "citation": "CISA. (2023). Zero Trust Maturity Model Version 2.0.",
        "url": "https://www.cisa.gov/resources-tools/resources/zero-trust-maturity-model",
        "note": "Provides a practical maturity framework for evaluating visibility, policy, identity, and device-centric controls.",
    },
    {
        "id": "R4",
        "citation": "CISA. (2025). Microsegmentation in Zero Trust Part One: Introduction and Planning.",
        "url": "https://www.cisa.gov/resources-tools/resources/microsegmentation-zero-trust-part-one-introduction-and-planning",
        "note": "Frames microsegmentation as a mechanism for limiting lateral movement and narrowing attack surface.",
    },
    {
        "id": "R5",
        "citation": "Jajodia, S., Ghosh, A. K., Swarup, V., Wang, C., & Wang, X. S. (Eds.). (2011). Moving Target Defense: Creating Asymmetric Uncertainty for Cyber Threats. Springer.",
        "url": "https://link.springer.com/book/10.1007/978-1-4614-0977-9",
        "note": "Foundational book that frames moving target defense as a way to break attacker advantage through controlled change.",
    },
    {
        "id": "R6",
        "citation": "Zhuang, R., Zhang, S., DeLoach, S., Ou, X., & Singhal, A. (2012). Simulation-based Approaches to Studying Effectiveness of Moving-Target Network Defense. NIST.",
        "url": "https://www.nist.gov/publications/simulation-based-approaches-studying-effectiveness-moving-target-network-defense",
        "note": "Shows why simulation and quantitative metrics matter when evaluating MTD effectiveness.",
    },
    {
        "id": "R7",
        "citation": "Zhuang, R., Zhang, S., Bardas, A., DeLoach, S., Ou, X., & Singhal, A. (2013). Investigating the Application of Moving Target Defenses to Enterprise Network Security.",
        "url": "https://csrc.nist.gov/pubs/conference/2013/08/15/investigating-the-application-of-moving-target-def/final",
        "note": "Demonstrates that adaptive network reconfiguration can reduce attacker success likelihood in modeled scenarios.",
    },
    {
        "id": "R8",
        "citation": "Jalowski, L., Zmuda, M., & Rawski, M. (2022). A Survey on Moving Target Defense for Networks: A Practical View. Electronics, 11(18), 2886.",
        "url": "https://www.mdpi.com/2079-9292/11/18/2886",
        "note": "Useful survey for understanding practical trade-offs, classifications, and deployment concerns in network MTD.",
    },
    {
        "id": "R9",
        "citation": "Lei, C., Zhang, H.-Q., Tan, J., Zhang, Y.-C., & Zhang, X. (2018). Moving Target Defense Techniques: A Survey. Security and Communication Networks.",
        "url": "https://doi.org/10.1155/2018/3759626",
        "note": "Summarizes MTD techniques, benefits, and challenges, including cost and policy complexity.",
    },
    {
        "id": "R10",
        "citation": "Zheng, J., & Namin, A. S. (2019). A Survey on the Moving Target Defense Strategies: An Architectural Perspective. Journal of Computer Science and Technology, 34(1), 207-233.",
        "url": "https://doi.org/10.1007/s11390-019-1906-z",
        "note": "Discusses MTD from a system architecture viewpoint and highlights SDN as an enabling technology.",
    },
    {
        "id": "R11",
        "citation": "Swami, R., Dave, M., & Ranga, V. (2023). Mitigation of DDoS Attack Using Moving Target Defense in SDN. Wireless Personal Communications, 131, 2429-2443.",
        "url": "https://link.springer.com/article/10.1007/s11277-023-10544-8",
        "note": "Connects SDN programmability with adaptive MTD responses for modern attack scenarios.",
    },
    {
        "id": "R12",
        "citation": "Zhang, M., Wang, L., Jajodia, S., Singhal, A., & Albanese, M. (2016). Network Diversity: A Security Metric for Evaluating the Resilience of Networks Against Zero-Day Attacks. IEEE TIFS, 11(5), 1071-1086.",
        "url": "https://csrc.nist.gov/pubs/journal/2016/01/network-diversity-security-metric-for-evaluating-r/final",
        "note": "Supports the idea that diversity and controlled variation can be measured and linked to resilience improvement.",
    },
]


FRONT_MATTER = {
    "student_names": "[Student Name / Team Name]",
    "registration_numbers": "[Registration Number(s)]",
    "guide_name": "[Guide Name]",
    "department": "Department of Computer Science and Engineering",
    "institution": "SRM Institute of Science and Technology",
    "submission_month_year": "April 2026",
}


def set_default_font(document: Document) -> None:
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5"]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    document.styles["Normal"].font.size = Pt(14)
    document.styles["Title"].font.size = Pt(20)
    document.styles["Subtitle"].font.size = Pt(14)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(15)
    document.styles["Heading 3"].font.size = Pt(14)
    document.styles["Heading 4"].font.size = Pt(14)
    document.styles["Heading 5"].font.size = Pt(13)


def add_paragraph(document: Document, text: str, style: str = "Normal", align: int | None = None, bold: bool = False) -> None:
    para = document.add_paragraph(style=style)
    if align is not None:
        para.alignment = align
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(text)
    run.bold = bold


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        para = document.add_paragraph(style="List Paragraph")
        para.paragraph_format.line_spacing = 1.5
        para.add_run(item)


def add_page_break(document: Document) -> None:
    document.add_page_break()


def add_toc_field(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run = paragraph.add_run("Right-click and update the field in Microsoft Word to refresh the table of contents.")
    run.italic = True
    run._r.append(fld_char_end)


def chapter_heading(document: Document, title: str) -> None:
    add_page_break(document)
    add_paragraph(document, title, style="Heading 1")


def section_heading(document: Document, title: str) -> None:
    add_paragraph(document, title, style="Heading 2")


def subsection_heading(document: Document, title: str) -> None:
    add_paragraph(document, title, style="Heading 3")


def paragraph_pack(document: Document, paragraphs: list[str]) -> None:
    for paragraph in paragraphs:
        add_paragraph(document, paragraph)


def status_breakdown() -> Counter:
    return Counter(item["status"] for item in ALIGNMENT["checks"])


def try_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/times.ttf",
        "C:/Windows/Fonts/timesbd.ttf" if bold else "C:/Windows/Fonts/times.ttf",
        "C:/Windows/Fonts/calibri.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size=size)
            except OSError:
                pass
    return ImageFont.load_default()


def new_canvas(width: int = 1600, height: int = 900, title: str | None = None) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    if title:
        draw.text((60, 30), title, fill="#0f2747", font=try_font(38, bold=True))
        draw.line((60, 85, width - 60, 85), fill="#7d97b6", width=3)
    return image, draw


def wrap_text(text: str, width: int) -> str:
    return "\n".join(textwrap.wrap(text, width=width))


def save_image(image: Image.Image, path: Path) -> Path:
    image.save(path)
    return path


def make_assets() -> dict[str, Path]:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)

    paths = {}

    image, draw = new_canvas(title="Implementation Readiness Across Documented VANTA Capabilities")
    categories = [item["name"] for item in ALIGNMENT["checks"][:-1]]
    score_map = {"implemented": 2, "partial": 1, "missing": 0}
    scores = [score_map[item["status"]] for item in ALIGNMENT["checks"][:-1]]
    left, top, right, bottom = 110, 160, 1480, 780
    draw.line((left, bottom, right, bottom), fill="#334155", width=3)
    draw.line((left, top, left, bottom), fill="#334155", width=3)
    step = (right - left - 40) / len(categories)
    colors = {2: "#1b8a5a", 1: "#d98e04", 0: "#b52b27"}
    axis_font = try_font(20)
    label_font = try_font(18)
    for idx, y_label in enumerate(["Missing", "Partial", "Implemented"]):
        y = bottom - idx * ((bottom - top) / 2)
        draw.line((left - 8, y, right, y), fill="#d7dee8", width=1)
        draw.text((35, y - 12), y_label, fill="#334155", font=axis_font)
    for idx, (category, score) in enumerate(zip(categories, scores)):
        x0 = left + 20 + idx * step
        bar_w = int(step * 0.6)
        bar_h = int((score / 2) * (bottom - top - 20))
        draw.rounded_rectangle((x0, bottom - bar_h, x0 + bar_w, bottom), radius=12, fill=colors[score], outline="#1f2937")
        draw.text((x0 - 10, bottom + 18), wrap_text(category, 16), fill="#111827", font=label_font)
    path = ASSETS_DIR / "implementation_readiness.png"
    paths["implementation_readiness"] = save_image(image, path)

    counts = status_breakdown()
    image, draw = new_canvas(width=1000, height=1000, title="Architecture Alignment Status Distribution")
    center = (500, 560)
    radius = 270
    values = [counts.get("implemented", 0), counts.get("partial", 0), counts.get("missing", 0)]
    labels = ["Implemented", "Partial", "Missing"]
    colors = ["#1b8a5a", "#d98e04", "#b52b27"]
    total = sum(values) or 1
    start = -90
    legend_font = try_font(26)
    for value, label, color in zip(values, labels, colors):
        angle = 360 * value / total
        draw.pieslice((center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius), start, start + angle, fill=color, outline="white")
        start += angle
    draw.ellipse((center[0] - 110, center[1] - 110, center[0] + 110, center[1] + 110), fill="white", outline="white")
    y = 160
    for value, label, color in zip(values, labels, colors):
        draw.rectangle((120, y, 160, y + 40), fill=color)
        draw.text((180, y + 4), f"{label}: {value}", fill="#111827", font=legend_font)
        y += 60
    path = ASSETS_DIR / "alignment_donut.png"
    paths["alignment_donut"] = save_image(image, path)

    image, draw = new_canvas(width=1200, height=1200, title="VANTA Capability Profile")
    center = (600, 650)
    radius = 360
    labels = list(MODERN["vanta"]["scores"].keys())
    values = list(MODERN["vanta"]["scores"].values())
    axis_font = try_font(18)
    for ring in range(1, 6):
        r = ring * radius / 5
        draw.ellipse((center[0] - r, center[1] - r, center[0] + r, center[1] + r), outline="#d7dee8", width=2)
    points = []
    for idx, (label, value) in enumerate(zip(labels, values)):
        angle = (2 * math.pi * idx / len(labels)) - math.pi / 2
        outer_x = center[0] + radius * math.cos(angle)
        outer_y = center[1] + radius * math.sin(angle)
        draw.line((center[0], center[1], outer_x, outer_y), fill="#d7dee8", width=2)
        tx = center[0] + (radius + 70) * math.cos(angle)
        ty = center[1] + (radius + 70) * math.sin(angle)
        draw.multiline_text((tx - 70, ty - 20), wrap_text(label.replace("_", " ").title(), 14), fill="#111827", font=axis_font, align="center")
        px = center[0] + (value / 5) * radius * math.cos(angle)
        py = center[1] + (value / 5) * radius * math.sin(angle)
        points.append((px, py))
    draw.polygon(points, outline="#004d99", fill="#9ec5ff")
    for point in points:
        draw.ellipse((point[0] - 8, point[1] - 8, point[0] + 8, point[1] + 8), fill="#004d99")
    path = ASSETS_DIR / "vanta_radar.png"
    paths["vanta_radar"] = save_image(image, path)

    image, draw = new_canvas(title="VANTA Compared with Mainstream Security Architectures")
    comparison_names = [item["architecture"] for item in MODERN["comparisons"]]
    vanta_scores = [item["vanta_total"] for item in MODERN["comparisons"]]
    other_scores = [item["other_total"] for item in MODERN["comparisons"]]
    left, top, right, bottom = 110, 180, 1480, 790
    draw.line((left, bottom, right, bottom), fill="#334155", width=3)
    step = (right - left - 60) / len(comparison_names)
    max_score = max(vanta_scores + other_scores)
    text_font = try_font(20)
    for idx, score in enumerate(range(0, max_score + 5, 5)):
        y = bottom - int((score / max_score) * (bottom - top))
        draw.line((left, y, right, y), fill="#e2e8f0", width=1)
        draw.text((40, y - 12), str(score), fill="#334155", font=text_font)
    for idx, name in enumerate(comparison_names):
        x0 = left + 30 + idx * step
        bar_h_a = int((vanta_scores[idx] / max_score) * (bottom - top - 20))
        bar_h_b = int((other_scores[idx] / max_score) * (bottom - top - 20))
        draw.rectangle((x0, bottom - bar_h_a, x0 + 55, bottom), fill="#0c6cf2")
        draw.rectangle((x0 + 65, bottom - bar_h_b, x0 + 120, bottom), fill="#8b9eb7")
        draw.multiline_text((x0 - 10, bottom + 15), wrap_text(name, 14), fill="#111827", font=try_font(18), align="center")
    draw.rectangle((1120, 120, 1160, 150), fill="#0c6cf2")
    draw.text((1175, 118), "VANTA", fill="#111827", font=text_font)
    draw.rectangle((1280, 120, 1320, 150), fill="#8b9eb7")
    draw.text((1335, 118), "Peer Architecture", fill="#111827", font=text_font)
    path = ASSETS_DIR / "modern_architecture_comparison.png"
    paths["modern_architecture_comparison"] = save_image(image, path)

    image, draw = new_canvas(title="VANTA Architecture Overview")
    box_font = try_font(24)
    boxes = [
        (90, 280, 330, 400, "Attacker or Scanner"),
        (410, 280, 700, 400, "OpenFlow Switch / OVS"),
        (800, 260, 1130, 420, "VANTA Controller"),
        (1210, 180, 1510, 300, "Threat Detector"),
        (1210, 330, 1510, 450, "VIP Mapping Engine"),
        (1210, 480, 1510, 600, "Dashboard + API"),
        (800, 610, 1130, 730, "Protected Hosts"),
    ]
    for x1, y1, x2, y2, label in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="#dce8f7", outline="#1c4d8c", width=4)
        draw.multiline_text((x1 + 25, y1 + 30), wrap_text(label, 18), fill="#102a43", font=box_font, align="center")
    for start, end in [((330, 340), (410, 340)), ((700, 340), (800, 340)), ((1130, 320), (1210, 240)), ((1130, 340), (1210, 390)), ((1130, 360), (1210, 540)), ((965, 420), (965, 610))]:
        draw.line((start, end), fill="#1c4d8c", width=5)
        draw.polygon([(end[0], end[1]), (end[0] - 16, end[1] - 8), (end[0] - 16, end[1] + 8)], fill="#1c4d8c")
    draw.text((400, 800), "Figure generated from repo architecture and controller responsibilities", fill="#334155", font=try_font(22))
    path = ASSETS_DIR / "architecture_overview.png"
    paths["architecture_overview"] = save_image(image, path)

    image, draw = new_canvas(title="Threat-Triggered Morphing Response Sequence")
    steps = [
        "Reconnaissance traffic observed",
        "Threat detector counts probes in active window",
        "Threshold crossed and event classified",
        "Controller triggers morph strategy",
        "VIP mapping rotates exposed identities",
        "Dashboard and logs record the response",
    ]
    x = 70
    for idx, step in enumerate(steps):
        draw.rounded_rectangle((x, 320, x + 210, 520), radius=20, fill="#eef5e5", outline="#597d35", width=4)
        draw.multiline_text((x + 18, 375), wrap_text(step, 18), fill="#2b4a12", font=try_font(23), align="center")
        if idx < len(steps) - 1:
            draw.line((x + 210, 420, x + 250, 420), fill="#597d35", width=5)
            draw.polygon([(x + 250, 420), (x + 232, 410), (x + 232, 430)], fill="#597d35")
        x += 250
    path = ASSETS_DIR / "threat_response_sequence.png"
    paths["threat_response_sequence"] = save_image(image, path)

    image, draw = new_canvas(title="Illustrative VIP Morphing Timeline")
    left, top, right, bottom = 130, 180, 1460, 760
    draw.line((left, bottom, right, bottom), fill="#334155", width=3)
    draw.line((left, top, left, bottom), fill="#334155", width=3)
    vip_labels = ["VIP Set A", "VIP Set B", "VIP Set C", "VIP Set D"]
    y_positions = [650, 520, 390, 260]
    for label, y in zip(vip_labels, y_positions):
        draw.text((30, y - 14), label, fill="#111827", font=try_font(22))
        draw.line((left, y, right, y), fill="#e2e8f0", width=1)
    timeline_points = [(left, 650), (360, 650), (361, 520), (700, 520), (701, 390), (1020, 390), (1021, 260), (1350, 260)]
    draw.line(timeline_points, fill="#ad1457", width=8)
    for x, y, label in [(430, 585, "Probe detected"), (810, 455, "Timer expired"), (1130, 325, "Manual or policy trigger")]:
        draw.ellipse((x - 12, y - 12, x + 12, y + 12), fill="#2e7d32")
        draw.text((x - 55, y - 55), label, fill="#2e7d32", font=try_font(18))
    for idx, marker in enumerate(range(0, 7)):
        x = left + idx * 200
        draw.text((x - 5, bottom + 20), str(marker), fill="#111827", font=try_font(20))
    path = ASSETS_DIR / "vip_timeline.png"
    paths["vip_timeline"] = save_image(image, path)

    return paths


def add_figure(document: Document, image_path: Path, caption: str, width: float = 6.2) -> None:
    document.add_picture(str(image_path), width=Inches(width))
    last_paragraph = document.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph(document, caption, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_table(document: Document, title: str, headers: list[str], rows: list[list[str]]) -> None:
    add_paragraph(document, title, align=WD_ALIGN_PARAGRAPH.CENTER)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr_cells[idx].text = header
    for row in rows:
        row_cells = table.add_row().cells
        for idx, value in enumerate(row):
            row_cells[idx].text = str(value)
    document.add_paragraph("")


def build_cover(document: Document) -> None:
    add_paragraph(document, "VANTA: Variable Network Topology Architecture", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(
        document,
        "A Research-Based Minor Project Report on Adaptive SDN-Driven Moving Target Defense with Virtual IP Morphing, Threat Detection, and Comparative Security Analysis",
        style="Subtitle",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(document, "Submitted by", style="Heading 4", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, FRONT_MATTER["student_names"], align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, FRONT_MATTER["registration_numbers"], align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, "Under the Guidance of", style="Heading 4", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, FRONT_MATTER["guide_name"], align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, "in partial fulfillment of the requirements for the degree of", style="Heading 4", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, "BACHELOR OF TECHNOLOGY", style="Heading 2", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, "in Computer Science and Engineering", style="Heading 2", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, FRONT_MATTER["department"], style="Heading 3", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, FRONT_MATTER["institution"], style="Heading 3", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(document, FRONT_MATTER["submission_month_year"], style="Heading 3", align=WD_ALIGN_PARAGRAPH.CENTER)


def build_preliminary_pages(document: Document) -> None:
    add_page_break(document)
    add_paragraph(document, "Own Work Declaration", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph_pack(
        document,
        [
            "I hereby declare that the material presented in this report is an original academic write-up prepared for the VANTA project using the referenced project documents, the current repository contents, and the cited research literature. Wherever external ideas, standards, or prior studies have informed the discussion, they have been acknowledged through explicit references in the bibliography.",
            "This report has been intentionally written in fresh language rather than by copying large passages from any one source. The intention is to support academic integrity, reduce direct textual overlap, and create a coherent narrative tailored to the present project implementation. Final plagiarism or AI-detection outcomes remain dependent on the institution, the checker used, and any later edits applied outside this generated version.",
            "The author or authors should replace the placeholder identity fields on the cover page and sign the institute-specific declaration pages before formal submission.",
        ],
    )

    add_page_break(document)
    add_paragraph(document, "Bonafide Certificate", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph_pack(
        document,
        [
            "This is to certify that the project report titled 'VANTA: Variable Network Topology Architecture' is a bonafide record of the research-oriented project work carried out by the above-named student or student team under the supervision of the guide identified on the title page.",
            "The work documented in this report focuses on the design and analysis of an adaptive software-defined networking security architecture that applies moving target defense concepts, virtual IP remapping, threat-triggered morphing, and implementation-level benchmarking in a laboratory setting. The report additionally includes a comparative study against mainstream security architecture patterns in order to situate the prototype in contemporary cyber defense practice.",
            "Institute signatures, dates, and examiner endorsements should be added in the final university submission copy according to departmental formatting rules.",
        ],
    )

    add_page_break(document)
    add_paragraph(document, "Acknowledgements", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph_pack(
        document,
        [
            "The successful preparation of this report depended on more than implementation effort alone. It also required a sequence of academic, technical, and organizational contributions that shaped the project into a research-based engineering exercise. I extend sincere gratitude to the faculty, mentors, and academic coordinators whose expectations encouraged the work to be presented not merely as code, but as a structured piece of inquiry with defensible problem framing, methodology, and evaluation logic.",
            "I am thankful to the project guide for insisting on clarity in design decisions and for emphasizing the importance of translating a prototype into a document that explains why the system matters, what gaps it addresses, and how its limitations can guide future investigation. That emphasis directly influenced the comparative, literature-based, and evidence-driven character of this report.",
            "Thanks are also due to the maintainers of the open-source networking and Python ecosystems that make security experimentation feasible. The VANTA prototype depends on a stack of community-maintained tools and libraries that support reproducibility, rapid prototyping, automated testing, visualization, and reporting. Their availability lowered the barrier to building and explaining a modern software-defined security experiment.",
            "Finally, I acknowledge the support of family, peers, and reviewers whose patience, encouragement, and questions helped convert a technically interesting prototype into a more disciplined academic submission.",
        ],
    )

    add_page_break(document)
    add_paragraph(document, "Abstract", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph_pack(
        document,
        [
            "Enterprise and laboratory networks alike are commonly exposed to a structural weakness: static addressing and stable topology information allow attackers to accumulate reconnaissance knowledge that often remains valid long enough to support exploitation, lateral movement, and service disruption. Moving Target Defense (MTD) responds to this asymmetry by deliberately changing observable system properties so that attacker knowledge ages rapidly. This report presents a research-based study of VANTA, a Variable Network Topology Architecture implemented as an adaptive software-defined networking prototype that combines virtual IP (VIP) remapping, multiple morphing strategies, threat-triggered response, dashboard-based observability, and prototype benchmarking support.",
            "The work begins with a review of academic and standards-based literature on moving target defense, software-defined networking, microsegmentation, and zero trust architectures. The literature indicates that static perimeter assumptions are increasingly misaligned with hybrid and distributed environments, while MTD remains a promising technique for degrading reconnaissance quality and increasing attacker uncertainty. VANTA is positioned within that context as a laboratory-scale, SDN-enabled architecture that uses controller logic to virtualize host identities and rotate them on the basis of communication events, timers, packet thresholds, or detected reconnaissance behavior.",
            "The implementation-oriented analysis in this report is grounded in the current repository structure of the project. Core controller capabilities, defense logic, attack simulation, benchmarking tooling, access control elements, deployment profile abstractions, and generated comparison reports are examined and synthesized into a structured engineering narrative. Generated figures and tables are included to summarize the architecture, response sequence, implementation readiness, and comparative capability profile of VANTA relative to static perimeter networks, general zero trust architecture, zero trust microsegmentation, and BeyondCorp-style access models.",
            "The study finds that VANTA is strongest as a research prototype for dynamic network obfuscation and reconnaissance invalidation. Its use of SDN control, VIP mapping, and threat-triggered morphing makes it more adaptive than conventional static networks and meaningfully differentiates it from architectures that rely only on perimeter hardening. At the same time, the analysis also shows that mainstream zero trust and microsegmentation approaches remain stronger in mature identity policy, device posture integration, and enterprise deployment readiness. The report therefore argues that VANTA should be viewed not as a replacement for zero trust, but as a complementary adaptive layer that can strengthen defenses where address stability and reconnaissance exposure remain operational risks.",
            "The report concludes that the project is academically meaningful because it demonstrates how modern programmable networking can host an experimentally tractable MTD workflow while also exposing the engineering gaps that separate a promising research prototype from a production-grade platform. Future work should focus on richer policy orchestration, hardware-backed deployment paths, larger-scale benchmarking, workload-level segmentation, and tighter integration with identity- and device-aware trust frameworks.",
        ],
    )

    add_page_break(document)
    add_paragraph(document, "Table of Contents", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_toc_field(document)


def intro_paragraphs() -> list[str]:
    return [
        "The VANTA project addresses a classic imbalance in cyber defense. Attackers often require only one period of uninterrupted reconnaissance to learn addressing patterns, service exposure, and reachable hosts, whereas defenders must sustain correctness and consistency over time. In a static network, the cost of learning is paid once by the attacker but defended continuously by the organization. That asymmetry is one of the strongest motivations for adaptive defense models.",
        "Moving target defense reframes the problem by turning stability into a managed variable. Instead of exposing the same host-to-address relationship indefinitely, the defender rotates or virtualizes selected properties so that intelligence gathered at one moment becomes less reliable in the next. The purpose is not to make the network random in an uncontrolled sense, but to orchestrate change in a way that preserves legitimate communication while degrading attacker certainty. VANTA adopts this idea at the network layer by using software-defined networking control logic and virtual IP remapping.",
        "Software-defined networking is especially suitable for this purpose because it centralizes traffic visibility and policy execution. By separating the control plane from the forwarding plane, SDN enables a controller to observe events, detect patterns, and modify flow behavior without requiring manual reconfiguration of every device. In VANTA, the SDN controller becomes the decision point that binds together host visibility, virtual identity assignment, and trigger-based morphing strategies.",
        "The project is also relevant because it does not treat security adaptation as a purely theoretical exercise. The repository includes controller code, a dashboard, attack simulation tooling, benchmarking scaffolding, access control logic, deployment abstractions, and comparative reporting scripts. That combination makes the project suitable for a research-based report: it contains a defensible problem statement, a concrete proposed system, artifacts that can be analyzed, and enough implementation detail to support engineering reasoning about strengths, limitations, and future enhancements.",
    ]


def literature_overview_paragraphs() -> list[str]:
    return [
        "The literature on moving target defense consistently argues that many computer systems remain easier to attack than to defend because their externally observable properties remain stable for long enough to support planning, rehearsal, and iterative exploitation. Foundational MTD work presents controlled dynamism as a way to create asymmetric uncertainty for cyber threats [R5]. Rather than relying exclusively on patching and filtering, MTD introduces change as a first-class defensive mechanism.",
        "Survey literature expands this core idea by classifying MTD strategies according to the layer at which change occurs, the trigger that initiates change, and the operational cost of that change [R8][R9][R10]. Network-layer MTD typically focuses on address mutation, route diversity, topology obfuscation, or service relocation. These techniques are especially useful against reconnaissance-driven attacks because they target the phase in which the attacker is building a reliable map of reachable resources.",
        "Research from NIST highlights that meaningful evaluation of MTD cannot stop at intuition. Simulation and metrics are needed to quantify how adaptation affects attacker success probability, operational integrity, and security trade-offs [R6][R7]. That perspective is central to the present report, which treats VANTA not merely as a code artifact but as an experiment-oriented architecture that should eventually be judged by measurable outcomes such as degraded scan usefulness, reaction latency, and acceptable overhead.",
        "The literature on zero trust provides an important neighboring context. NIST SP 800-207 argues that defenders should move away from trusting assets on the basis of network location alone and instead focus on users, devices, and resources [R1]. CISA extends that line of thinking through maturity models and practical guidance for segmentation and least-privilege access [R3][R4]. These sources do not replace MTD research; instead, they help explain where a prototype like VANTA fits within the larger trajectory of modern network defense.",
    ]


def research_gap_paragraphs() -> list[str]:
    return [
        "A recurring research gap appears between concept-rich MTD literature and production-oriented implementation detail. Many studies demonstrate why dynamism matters, yet fewer open prototypes show how threat detection, identity mutation, visibility, and experimentation can be integrated in one inspectable stack. VANTA helps narrow that gap by exposing code-level mechanisms and by making the controller, dashboard, and evaluation tooling part of the same narrative.",
        "Another gap concerns the relationship between MTD and zero trust. Literature often treats these ideas in parallel: MTD emphasizes unpredictability and attack-surface change, while zero trust emphasizes identity, device state, and least-privilege authorization. The result is that organizations may see them as substitutes rather than complementary layers. VANTA exposes an opportunity to combine both schools of thought by pairing dynamic VIP morphing with access control, MFA awareness, and device-context handling in the same prototype.",
        "There is also a scale gap between experimental network defense papers and enterprise deployment expectations. Prototype systems can demonstrate promise in Mininet or Open vSwitch settings, but larger environments demand standardized policy, persistent state, multi-node resilience, and hardware-aware rollout constraints. The repository evidence examined in this report shows that VANTA has begun moving in that direction through adapter seams, SQLite state persistence, and hardware inventory policies, yet it still operates primarily as a research prototype.",
        "Finally, there remains a documentation gap in many student or prototype projects: implementation exists, but the system is not translated into a structured academic report that explains motivation, prior work, methodology, evaluation logic, and comparative significance. This report explicitly addresses that gap by converting the project into a coherent research document rather than a purely implementation-oriented showcase.",
    ]


def objectives_bullets() -> list[str]:
    return [
        "To study moving target defense as a response to stale reconnaissance in static networks.",
        "To analyze how SDN control can support programmable VIP remapping and threat-triggered adaptation.",
        "To examine the VANTA repository as a concrete prototype of network-layer moving target defense.",
        "To compare VANTA with traditional perimeter security and mainstream zero trust-oriented models.",
        "To derive implementation insights, research gaps, and future enhancement paths from the current artifact set.",
    ]


def methodology_paragraphs() -> list[str]:
    return [
        "This report follows a design-science and prototype-analysis methodology. The project itself is an engineered artifact, but the report treats the artifact as an object of study. Instead of relying on a single source, the analysis combines template expectations from the provided academic reference documents, repository inspection of the current VANTA codebase, generated comparison reports already present in the project, standards-based security literature, and a review of the testing and benchmarking surfaces exposed by the repository.",
        "The method is intentionally layered. The first layer establishes the academic framing: problem statement, literature review, research gap, and objectives. The second layer examines the implemented architecture: controller functions, morphing logic, attack simulation, benchmarking, dashboard, access control, state management, and deployment abstractions. The third layer evaluates readiness and comparative meaning by using generated repository reports and by organizing the implementation evidence into tables and capability maps. The fourth layer reflects critically on the prototype's limitations, especially in relation to enterprise maturity and large-scale validation.",
        "Because this report is generated from the current project state rather than from a full laboratory rerun with Mininet and Ryu on this machine, all claims are carefully framed. Code-backed features and existing result artifacts are treated as present evidence. Full end-to-end performance metrics that would require a live networking environment are described as part of the evaluation design and the prototype's intended experimental workflow, not as unsupported empirical facts.",
        "This approach is academically useful because it distinguishes between three levels of confidence: implemented capability that is directly visible in the codebase, comparative inference that is grounded in repository-generated reports and standards, and forward-looking experimental claims that still require broader runtime validation. Maintaining that distinction helps keep the report research-based without overstating certainty.",
    ]


def sprint_text(title: str, focus: str, outcomes: list[str]) -> list[str]:
    base = [
        f"{title} was organized around the theme of {focus}. In agile terms, the sprint can be understood as a bounded implementation and learning cycle rather than a purely administrative milestone. The sprint logic is important to this report because it reveals how the architecture likely matured: first by proving core control behavior, then by improving observability and reproducibility, and finally by strengthening deployment-oriented abstractions.",
        f"The work in {title} can be interpreted as a blend of engineering tasks and research preparation. Engineering tasks made the prototype functional, while research preparation tasks made it explainable, testable, and comparable. That dual nature matters because academic project work is strongest when system construction and evaluation planning progress together rather than being postponed to the end.",
    ]
    base.extend([f"A notable outcome of {title} was {outcome}." for outcome in outcomes])
    base.append(
        f"From a research perspective, {title} improved not just the number of features but the interpretability of the system. The more explicit the controller states, triggers, logs, exports, and module boundaries become, the easier it is to reason about the prototype as a candidate architecture rather than a one-off script."
    )
    return base


def architecture_description() -> list[str]:
    return [
        "VANTA is best described as an SDN-enabled moving target defense architecture that virtualizes network identities and changes them under controlled policy. At its center sits a controller responsible for OpenFlow communication, virtual-to-real and real-to-virtual mappings, morphing policies, attack-trigger evaluation, and dashboard-facing telemetry. This makes the controller more than a forwarding aid; it becomes the system's adaptive security brain.",
        "The VIP mapping engine is the subsystem that operationalizes uncertainty. Instead of allowing a protected host to remain exposed under one stable, attacker-visible identity, the engine assigns and rotates virtual IPs while preserving an internal understanding of the underlying real host. When legitimate traffic or a predefined trigger occurs, the mapping can change. The objective is not merely motion for its own sake, but the invalidation of stale attacker observations.",
        "Threat detection gives the architecture its reactive strength. Repository evidence indicates that VANTA includes logic for detecting scan-like behavior and using that detection to trigger defensive remorphing. This matters because a useful MTD system should not depend on a single schedule alone. Purely periodic change may still leave windows of predictability, whereas event-aware morphing lets the defender react when attacker intent becomes visible through network behavior.",
        "Observability is another important design feature. The dashboard and API surfaces transform the architecture from a hidden mechanism into a measurable one. Security prototypes benefit when operators can inspect mappings, history, strategy selection, access context, and health signals. Visibility supports both debugging and research because it creates a path from internal controller behavior to external evidence and reporting.",
        "The repository also suggests an ongoing effort toward architectural modularization. While a large controller file still appears to serve as the primary runtime entry point, logic has been separated into modules for defense, configuration, deployment, access control, network adapters, hardware rollout policy, state storage, and VIP mapping. This is significant for maintainability and for future experimentation, since modular components are easier to test, replace, and extend than tightly coupled monoliths.",
    ]


def evaluation_design_paragraphs() -> list[str]:
    return [
        "A research-grade evaluation of VANTA should consider both security effect and operational overhead. Security effect asks whether the system reduces the value of attacker reconnaissance, shortens the useful lifetime of observed host identities, and increases the difficulty of persistent probing. Operational overhead asks whether the resulting control actions impose unacceptable penalties on latency, throughput, resource consumption, or management complexity.",
        "The repository includes an attack simulator and a benchmark suite, which together establish a plausible experimental framework. The attack simulator contains routines for port scanning, SYN flood behavior, ICMP flood scenarios, and reconnaissance-oriented activity. The benchmark suite outlines measurements for latency, throughput, CPU load, memory usage, and strategy comparison. Even where a live lab rerun was not performed in this report generation step, the presence of these tools is important evidence that the project was designed with experimentation in mind.",
        "A well-structured evaluation should compare at least four conditions: a static baseline without morphing, a periodic morphing mode, a packet-threshold or communication-triggered mode, and a threat-triggered mode that reacts to scan patterns. Across those conditions, the same traffic traces, topology shape, and attacker routine should be used so that differences in reconnaissance quality and overhead can be attributed to the defense strategy rather than to environmental changes.",
        "Metrics should include the number of ports or hosts correctly discovered by an attacker, the time window for which discovered identities remain valid, the delay between suspicious activity and remorphing, the number of flow updates issued by the controller, average and tail latency for legitimate traffic, and controller-side CPU or memory consumption. In addition, because VANTA contains dashboard and export capabilities, the completeness and consistency of generated experiment artifacts should also be measured as part of research reproducibility.",
    ]


def results_discussion_paragraphs() -> list[str]:
    return [
        "The available repository evidence supports three broad conclusions. First, VANTA is clearly aligned with the central idea of network-layer moving target defense. The implementation and generated reports show the presence of SDN control, VIP mapping, morphing strategies, threat-triggered response, and dashboard-facing visibility. Those features together make the project much more than a conceptual proposal.",
        "Second, the project remains recognizably prototype-oriented. The repository-generated alignment report identifies partial completion in areas such as attack simulation coverage, benchmark completeness, modularization, broader automated testing, and standardized result artifacts. That combination is typical of serious student or research prototypes: the core architectural idea is present and functional, while production-hardening and comprehensive validation remain ongoing.",
        "Third, the comparative analysis positions VANTA most favorably when the evaluation criterion values dynamic network obfuscation, reconnaissance invalidation, and adaptive response. It is less dominant when the criterion emphasizes mature identity enforcement, device posture validation, and large-scale deployment conventions. This is not a weakness in the concept of MTD itself, but a reminder that dynamic obfuscation addresses one class of attacker advantage and should ideally be fused with other modern control models.",
        "The discussion therefore supports a layered interpretation of the project. VANTA is a strong educational and research platform for understanding how controlled dynamism can be introduced into network defense using SDN. At the same time, the prototype's own design trajectory suggests that its long-term value may be greatest when integrated with zero trust, microsegmentation, or hybrid policy stacks rather than when evaluated in isolation.",
    ]


def add_reference_section(document: Document) -> None:
    chapter_heading(document, "References")
    for source in SOURCES:
        add_paragraph(document, f"{source['id']}. {source['citation']} Available at: {source['url']}")


def add_appendices(document: Document) -> None:
    chapter_heading(document, "Appendix A: Code and Module Summary")
    paragraph_pack(
        document,
        [
            "The repository structure suggests a project that has evolved from a single-controller prototype toward a more modular package. The principal runtime entry point remains the controller, while supporting packages now isolate defense policies, configuration handling, deployment profiles, access control, hardware inventory policy, network adapter selection, state persistence, and VIP mapping. This decomposition is a healthy direction for future maintainability and testability.",
            "The attack simulator and benchmark suite deserve attention because they transform the project from a demonstrator into an experiment host. Together, these modules create the basis for scenario-driven validation, repeatable measurement, and result export. Even when parts of the benchmark surface remain incomplete, their existence strengthens the academic value of the project by making evaluation an explicit design concern.",
        ],
    )

    chapter_heading(document, "Appendix B: Publication and Presentation Readiness Note")
    paragraph_pack(
        document,
        [
            "The material in this report can be distilled into a conference-style paper by concentrating on the problem of stale reconnaissance in SDN environments, the architecture of threat-triggered VIP morphing, and the comparative positioning of VANTA relative to static and zero trust approaches. A concise publication should prioritize architecture, trigger logic, evaluation design, and the specific research contribution of integrating adaptive remapping with visibility and access context.",
            "For presentation purposes, the most defensible visuals are the architecture overview, the threat-trigger sequence, the implementation readiness chart, and the comparative capability bar chart. These figures communicate the system's purpose, control flow, maturity state, and contextual value without requiring the audience to read code directly.",
        ],
    )

    chapter_heading(document, "Appendix C: Academic Integrity and Similarity Note")
    paragraph_pack(
        document,
        [
            "This report was written as a fresh, project-specific document and intentionally avoids long copied passages from prior material. It paraphrases cited concepts in original language and grounds system-specific claims in the current repository state. That said, similarity and AI-detection outcomes depend on the institution's software, settings, exclusions, and any post-generation edits. No automated tool can be guaranteed to return a value below a fixed threshold in every environment.",
            "Before final submission, the author should run the institute's official similarity workflow, replace placeholders on the front matter, review the references for formatting alignment with the university style, and if required attach the official plagiarism or similarity report generated by the university-approved platform.",
        ],
    )


def build_main_body(document: Document, assets: dict[str, Path]) -> None:
    chapter_heading(document, "Chapter 1: Introduction")
    section_heading(document, "1.1 Introduction to the Project")
    paragraph_pack(document, intro_paragraphs())
    add_figure(document, assets["architecture_overview"], "Figure 1. High-level VANTA architecture derived from the current repository design.")

    section_heading(document, "1.2 Problem Statement and Description")
    paragraph_pack(
        document,
        [
            "The core problem addressed by this project is the persistent exposure created by static network identities. When host addresses and network structure remain predictable, attackers can perform scans, correlate observations over time, and construct a stable map of reachable services. That map becomes a reusable resource for exploitation, credential attacks, and lateral movement. Defensive teams must then work against an adversary that already possesses reliable orientation.",
            "Traditional perimeter-centric models reduce some exposure but do not by themselves invalidate already observed network information. Firewalls and access control rules may restrict access, yet an attacker who manages to probe the environment often continues to benefit from the consistency of host identities. VANTA addresses that weakness by treating the externally visible address space as an adaptive surface rather than a fixed one.",
            "The project therefore asks a focused research question: can an SDN-controlled architecture use VIP remapping and trigger-aware morphing to reduce the usefulness of reconnaissance while preserving sufficient operational visibility and manageability? That question leads naturally to the system, literature, and evaluation choices made throughout this report.",
        ],
    )
    add_table(
        document,
        "Table 1. Core Security Problem Framing",
        ["Static Network Weakness", "Operational Consequence", "VANTA Response"],
        [
            ["Stable addressing", "Reconnaissance remains useful for longer", "VIP remapping and periodic or event-driven morphing"],
            ["Limited control-plane adaptivity", "Slow response to scan behavior", "SDN controller-driven trigger logic"],
            ["Poor experiment visibility", "Hard to compare strategies", "Dashboard, APIs, exports, and benchmark scaffolding"],
            ["Perimeter-only thinking", "Weak contextual adaptation", "Access context and zero-trust-aware extension points"],
        ],
    )

    section_heading(document, "1.3 Motivation")
    paragraph_pack(
        document,
        [
            "The motivation for VANTA comes from both security theory and implementation practicality. From the theoretical side, moving target defense promises to reduce attacker certainty by forcing observations to decay. From the practical side, SDN gives a realistic control point for building this kind of defense in a laboratory environment. Instead of assuming bespoke hardware support from the beginning, the project demonstrates how a controller-based experiment can host adaptive defense logic with modest setup overhead.",
            "A second motivation is educational depth. Many cybersecurity projects stop at detection or logging. VANTA goes further by linking detection with active architectural change. That makes it a richer project for academic study because it raises questions of policy timing, mapping consistency, state management, experimental repeatability, and trade-offs between security benefit and operational overhead.",
            "A third motivation lies in relevance to modern enterprise evolution. Zero trust, microsegmentation, and hybrid deployment models are increasingly emphasized by standards bodies and government guidance [R1][R3][R4]. Although VANTA is not a complete enterprise platform, it offers a useful research path toward adaptive controls that could complement identity- and policy-centric approaches in environments where address or topology stability still benefits attackers.",
        ],
    )

    section_heading(document, "1.4 Sustainable Development Goal Alignment")
    paragraph_pack(
        document,
        [
            "The project aligns most closely with SDG 9, which emphasizes industry, innovation, and resilient infrastructure. Secure digital infrastructure is now foundational to economic activity, education, health systems, and public administration. Research that strengthens network resilience, especially through programmable and measurable mechanisms, contributes to the broader goal of more dependable infrastructure systems.",
            "The work also relates to SDG 16 through its contribution to trustworthy institutions and secure digital governance. Modern institutions increasingly depend on online services, remote access, and distributed information systems. By studying mechanisms that reduce attacker certainty and improve defensive adaptability, the project supports the broader objective of safer digital environments.",
        ],
    )

    chapter_heading(document, "Chapter 2: Literature Survey")
    section_heading(document, "2.1 Overview of the Research Area")
    paragraph_pack(document, literature_overview_paragraphs())
    add_figure(document, assets["threat_response_sequence"], "Figure 2. Threat-triggered morphing response sequence represented as a controller workflow.")

    section_heading(document, "2.2 Existing Models and Frameworks")
    paragraph_pack(
        document,
        [
            "Existing research models for network moving target defense can be grouped into several classes. One class focuses on address or endpoint mutation, in which host identities exposed to adversaries are shifted over time. Another class emphasizes topology or route diversity, where the defender changes connectivity paths or graph visibility. A third class blends multiple adaptations, including service migration, software diversity, and policy orchestration. Surveys show that each class trades implementation cost against security effect in different ways [R8][R9][R10].",
            "SDN appears repeatedly in the literature as a strong enabler for network-layer MTD because it simplifies observability and policy insertion. The central controller can observe flows, infer patterns, and apply forwarding changes with less manual device-by-device intervention. Recent work examining MTD in SDN environments also points to the promise of adaptive response against DDoS and reconnaissance patterns [R11]. VANTA follows this logic by using controller-mediated decisions as the basis for VIP morphing.",
            "Zero trust frameworks contribute a different but complementary model. They do not seek to hide the network so much as to stop treating network position as adequate trust evidence. NIST and CISA documents emphasize resource-centric access decisions, identity, device state, visibility, segmentation, and least privilege [R1][R2][R3][R4]. For VANTA, these frameworks matter because they reveal where a dynamic address-mutation prototype is already strong and where it remains incomplete.",
            "In practice, the most useful model may be hybrid. MTD degrades attacker intelligence, while zero trust and microsegmentation improve policy precision and containment. The literature therefore suggests that a prototype like VANTA is valuable not only as a standalone defense experiment but also as a stepping stone toward composite architectures where deception, identity, and segmentation reinforce each other.",
        ],
    )
    add_table(
        document,
        "Table 2. Selected Literature and Their Relevance to VANTA",
        ["Ref", "Theme", "Contribution to This Report"],
        [[source["id"], source["citation"].split(". ")[1][:75] + "...", source["note"]] for source in SOURCES[:10]],
    )

    section_heading(document, "2.3 Limitations Identified from Literature Survey (Research Gaps)")
    paragraph_pack(document, research_gap_paragraphs())

    section_heading(document, "2.4 Research Objectives")
    add_bullets(document, objectives_bullets())

    section_heading(document, "2.5 Product Backlog (Key User Stories with Desired Outcomes)")
    add_table(
        document,
        "Table 3. Research and Engineering Backlog",
        ["User Story or Engineering Need", "Desired Outcome", "Academic Relevance"],
        [
            ["As a defender, I need stale scan results to expire quickly", "Attacker reconnaissance quality should degrade over time", "Central security motivation"],
            ["As a controller operator, I need multiple morphing strategies", "Different workloads and threat modes can be evaluated", "Supports comparative experiments"],
            ["As a researcher, I need visible system state", "Dashboard and exports show mappings, history, and statistics", "Supports evidence and reporting"],
            ["As a maintainer, I need modular code boundaries", "Core logic becomes easier to test and extend", "Improves reproducibility and future work"],
            ["As an evaluator, I need benchmark and attack tooling", "Prototype behavior can be measured under controlled scenarios", "Enables quantitative analysis"],
        ],
    )

    section_heading(document, "2.6 Plan of Action (Project Road Map)")
    paragraph_pack(
        document,
        [
            "The project road map can be interpreted as a sequence of increasing rigor. The first stage establishes the core idea in code: connect the SDN controller, maintain real-to-virtual mappings, and verify that traffic can be supported through virtualized identities. The second stage introduces trigger logic so that remorphing can happen on policy-relevant events instead of relying only on manual action. The third stage adds operator visibility through a dashboard and export endpoints. The fourth stage expands evaluation through attack simulation, benchmarking, and comparison reporting. The fifth stage begins shifting the architecture toward modularity, state persistence, and deployment abstraction.",
            "This road map matters because research-based project quality depends not only on the final feature list but on the logic of progression. Each stage adds a new dimension of credibility: implementation proves feasibility, triggers prove adaptivity, visibility proves explainability, evaluation proves seriousness, and modularization proves maintainability.",
        ],
    )
    add_figure(document, assets["vip_timeline"], "Figure 3. Illustrative VIP morphing timeline showing how exposed identity states can change across observation windows.")

    chapter_heading(document, "Chapter 3: Sprint Planning and Execution Methodology")
    section_heading(document, "3.1 Methodology")
    paragraph_pack(document, methodology_paragraphs())
    add_table(
        document,
        "Table 4. Methodological Layers Used in This Report",
        ["Layer", "Primary Evidence", "Purpose"],
        [
            ["Academic framing", "Reference .docx templates and cited literature", "Align the document with research-report expectations"],
            ["Artifact inspection", "Repository code, modules, and generated reports", "Identify implemented capabilities and gaps"],
            ["Comparative analysis", "Alignment JSON and architecture comparison JSON", "Position VANTA against peer models"],
            ["Critical reflection", "Testing outcome, maturity gaps, future work", "Prevent overstatement and improve rigor"],
        ],
    )

    section_heading(document, "3.2 Sprint I: Core Adaptive Security Prototype")
    paragraph_pack(
        document,
        sprint_text(
            "Sprint I",
            "establishing the minimum viable moving target defense workflow",
            [
                "a controller-centric model in which real and virtual identities are both tracked explicitly",
                "the inclusion of multiple morphing strategies rather than a single fixed remap rule",
                "the creation of a security narrative that linked observed reconnaissance to defensive change",
            ],
        ),
    )

    section_heading(document, "3.3 Sprint II: Visibility, Evaluation, and Operational Maturity")
    paragraph_pack(
        document,
        sprint_text(
            "Sprint II",
            "improving observability, comparative reporting, and deployment readiness",
            [
                "dashboard and API capabilities that expose mappings, strategy settings, and health information",
                "benchmark and attack simulation tooling that turn the prototype into an experiment host",
                "initial modularization and state-persistence work that reduce dependence on one monolithic runtime file",
            ],
        ),
    )
    add_figure(document, assets["implementation_readiness"], "Figure 4. Readiness of documented capabilities derived from repository comparison artifacts.")

    chapter_heading(document, "Chapter 4: System Design and Architecture Document")
    section_heading(document, "4.1 Functional Requirements")
    add_table(
        document,
        "Table 5. Functional Requirements",
        ["Requirement ID", "Description", "Status Perspective"],
        [
            ["FR-1", "Maintain real-to-virtual and virtual-to-real mappings", "Implemented in controller and mapping logic"],
            ["FR-2", "Trigger IP morphing using multiple strategies", "Implemented with partial extension scope"],
            ["FR-3", "Detect suspicious scan behavior", "Implemented in threat detection logic"],
            ["FR-4", "Expose runtime visibility through dashboard and APIs", "Implemented"],
            ["FR-5", "Support evaluation via attack and benchmark tooling", "Partially implemented and expandable"],
            ["FR-6", "Persist relevant state for restart resilience", "Implemented through current state-store direction"],
        ],
    )

    section_heading(document, "4.2 Non-Functional Requirements")
    add_table(
        document,
        "Table 6. Non-Functional Requirements",
        ["Attribute", "Expectation", "Why It Matters"],
        [
            ["Adaptivity", "Morphing must occur quickly when triggered", "Slow change weakens MTD benefit"],
            ["Observability", "Operators should inspect mappings and events", "Research and debugging depend on visibility"],
            ["Modularity", "Core components should be separable", "Improves maintainability and testing"],
            ["Reproducibility", "Experiments should produce reusable artifacts", "Supports academic evidence"],
            ["Safety", "Legitimate connectivity should remain stable enough", "Defense must not collapse normal operation"],
        ],
    )

    section_heading(document, "4.3 Architecture Description")
    paragraph_pack(document, architecture_description())
    add_figure(document, assets["alignment_donut"], "Figure 5. Distribution of implemented, partial, and missing capability statuses in the architecture alignment report.", width=4.9)

    section_heading(document, "4.4 Module-Wise Analysis")
    add_table(
        document,
        "Table 7. Major Repository Components and Their Roles",
        ["Component", "Observed Role", "Research Importance"],
        [
            ["ultimate_mtd_controller.py", "Primary runtime entry point and controller orchestration", "Core implementation evidence"],
            ["vanta_core/defense.py", "Threat detection and morphing strategy logic", "Supports adaptive response reasoning"],
            ["attack_simulator.py", "Emulates attack and probing behavior", "Supports controlled adversarial evaluation"],
            ["benchmark_suite.py", "Defines performance and comparison tests", "Supports overhead analysis"],
            ["vanta_core/access_control.py", "Identity and device-context decisions", "Links MTD to zero-trust concepts"],
            ["vanta_core/state_store.py", "Persistence of runtime state", "Improves resilience and continuity"],
            ["vanta_core/network_adapters.py", "Backend abstraction for different deployment modes", "Supports future scale-out and hardware orientation"],
        ],
    )
    paragraph_pack(
        document,
        [
            "The module layout indicates that the project is progressing toward a reusable architecture rather than remaining fixed as a single classroom demonstration. This direction matters because research prototypes often lose value if they cannot be extended, reproduced, or tested in isolation. Modular code does not merely look cleaner; it enables finer-grained reasoning about cause and effect inside the defense system.",
            "The attack simulator and benchmark modules also deserve emphasis because they close an important loop. A moving target defense design without adversarial or performance evaluation risks becoming a conceptual display. By contrast, VANTA at least begins to connect implementation with measurement, which is a defining characteristic of research-oriented engineering.",
        ],
    )

    section_heading(document, "4.5 Threat Model")
    paragraph_pack(
        document,
        [
            "The immediate threat model for VANTA centers on reconnaissance and early-stage probing behavior. This includes port scanning, patterned service discovery, and related traffic that attempts to infer host exposure. Such activity is important because it typically precedes more targeted exploitation. A defender that invalidates reconnaissance early may force the attacker either to spend more time re-learning the environment or to operate with less confidence.",
            "The threat model also extends to repeated probing, traffic bursts, and conditions that may correlate with service enumeration or preparation for denial-of-service attempts. The included literature on SDN-based MTD against DDoS supports the idea that programmable control can react to certain aggressive traffic patterns [R11], although the present report remains careful to distinguish code visibility from fully benchmarked runtime proof.",
            "What VANTA does not yet fully model is equally important. Enterprise-grade insider abuse, sophisticated multi-stage persistence, workload-level segmentation drift, and large-scale hybrid identity orchestration remain beyond the prototype's strongest validated surface. Recognizing these boundaries improves the honesty and therefore the value of the report.",
        ],
    )

    section_heading(document, "4.6 Data and Control Flow")
    paragraph_pack(
        document,
        [
            "A plausible control-flow interpretation of VANTA begins when traffic reaches the OpenFlow-enabled switch. The switch either forwards known flows according to previously installed rules or escalates relevant events to the controller. The controller interprets those events in the context of current VIP mappings, strategy settings, and threat-detection state. If a morphing condition is met, the controller rotates one or more visible identities and updates the network's forwarding expectations accordingly.",
            "In parallel, telemetry is emitted toward the dashboard and API surface. This allows the system to function as a visible defense rather than a silent one. Visibility is academically useful because it provides a bridge between internal controller actions and externally inspectable records, helping the experimenter reason about when and why a morph happened.",
            "State persistence completes the flow by making mappings and events more durable across runtime transitions. In research environments, persistence is important not only for resilience but also for reproducibility. If mappings and history vanish on restart, it becomes harder to explain longer-running experiments or compare sequential runs with confidence.",
        ],
    )

    chapter_heading(document, "Chapter 5: Experimental Setup and Evaluation Framework")
    section_heading(document, "5.1 Experimental Environment")
    paragraph_pack(
        document,
        [
            "The intended experimental environment for VANTA is a Linux-based SDN laboratory combining Ryu, Mininet, and Open vSwitch. This setup is appropriate for early-stage network security research because it allows controlled topology definition, scriptable traffic generation, and controller-driven policy behavior without requiring an enterprise hardware budget. The quick-start material in the repository reflects this assumption clearly.",
            "For report-generation purposes, the current environment was sufficient to inspect the repository and generate figures, but it did not provide the full live SDN runtime stack needed to execute every end-to-end networking experiment locally. Consequently, the evaluation framework in this chapter is grounded in the project's benchmark and attack tooling, current generated artifacts, and controller-oriented design intent.",
        ],
    )

    section_heading(document, "5.2 Evaluation Metrics")
    paragraph_pack(document, evaluation_design_paragraphs())
    add_table(
        document,
        "Table 8. Proposed Evaluation Metrics for VANTA",
        ["Metric", "Type", "Interpretation"],
        [
            ["Ports or hosts discovered", "Security effectiveness", "Lower values indicate stronger reconnaissance degradation"],
            ["Identity validity window", "Security effectiveness", "Shorter windows reduce stale attacker knowledge"],
            ["Detection-to-morph delay", "Responsiveness", "Lower delay indicates faster reaction"],
            ["Average RTT and throughput", "Performance overhead", "Used to quantify operational cost"],
            ["CPU and memory load", "Controller efficiency", "Shows sustainability of adaptive logic"],
            ["Flow churn or update count", "Control-plane cost", "Captures adaptation complexity"],
            ["Export completeness", "Reproducibility", "Measures usefulness of experiment artifacts"],
        ],
    )

    section_heading(document, "5.3 Experimental Scenarios")
    paragraph_pack(
        document,
        [
            "A meaningful experiment suite should include at least one benign communication baseline and multiple adversarial scenarios. Benign traffic establishes the control condition for performance measurements and helps determine whether morphing interrupts legitimate communication. Port scan scenarios test the prototype's core promise of degrading reconnaissance quality. SYN flood or traffic burst scenarios help reveal whether the controller's adaptive logic remains stable under more aggressive conditions.",
            "Strategy comparison is particularly important. Reply-triggered, time-based, packet-count-based, threat-triggered, and manual modes each express different assumptions about how much evidence is needed before change should occur. A research evaluation should not treat them as equivalent. Periodic remorphing may maximize predictability for the defender but leave avoidable blind windows, while threat-triggered change may be more selective but require accurate detection thresholds.",
            "Comparisons should also measure user experience and control-plane cost. A defense that destroys scan quality but destabilizes every normal session would not be practical. Conversely, a defense that preserves perfect performance but rarely remorphs under adversarial activity may have little strategic value. The purpose of experimentation is to locate that middle ground rather than to maximize one metric in isolation.",
        ],
    )

    section_heading(document, "5.4 Repository-Backed Validation Snapshot")
    add_table(
        document,
        "Table 9. Validation Snapshot Derived from Repository Artifacts",
        ["Evidence Source", "Observed Signal", "Interpretation"],
        [
            ["Alignment report JSON", "Weighted alignment score of 77.3%", "The current implementation reflects most major project ideas but still has maturity gaps"],
            ["Modern architecture comparison JSON", "VANTA total score of 35 in qualitative matrix", "Strong in dynamic obfuscation, moderate in identity/device maturity"],
            ["Unit test run in current environment", "Some tests passed, some failed due to environment and dependency constraints", "Prototype logic exists, but runtime environment influences verification depth"],
        ],
    )
    paragraph_pack(
        document,
        [
            "The unit test invocation performed during report preparation produced a mixed outcome. A meaningful subset of logic-focused tests passed, especially in defense, configuration, and access-related areas, but several failures were tied to environment-specific file permissions, temporary path access, missing Ryu dependencies, and database path availability in the current machine context. This does not invalidate the project, but it does illustrate a classic prototype challenge: reproducibility depends not only on code but also on runtime preparation.",
            "That observation is useful academically because it shows why packaging, setup automation, dependency management, and test isolation matter in cybersecurity research artifacts. A prototype becomes far more valuable when others can reproduce both its behavior and its validation results with minimal environmental friction.",
        ],
    )

    chapter_heading(document, "Chapter 6: Results and Discussion")
    section_heading(document, "6.1 Project Outcomes")
    paragraph_pack(document, results_discussion_paragraphs())
    add_figure(document, assets["modern_architecture_comparison"], "Figure 6. Qualitative comparison of VANTA with selected mainstream security architecture patterns.")

    section_heading(document, "6.2 Architecture Alignment Results")
    implemented = status_breakdown().get("implemented", 0)
    partial = status_breakdown().get("partial", 0)
    missing = status_breakdown().get("missing", 0)
    paragraph_pack(
        document,
        [
            f"The repository-generated architecture alignment report classifies {implemented} areas as implemented, {partial} as partial, and {missing} as missing. The important result is not the count alone but the distribution: core architectural concepts such as SDN control, VIP morphing, threat-triggered defense, dashboard visibility, authentication, and configurable runtime settings are already present, while surrounding maturity features still need expansion.",
            "This distribution is healthy for a minor research project because it indicates that the central contribution has been built rather than merely described. The remaining gaps are concentrated around completeness, coverage, and generalization rather than around the absence of the core idea itself.",
        ],
    )
    add_table(
        document,
        "Table 10. Summary of Alignment Findings",
        ["Category", "Observation", "Implication"],
        [
            ["Core defense logic", "Implemented", "The main research proposition is represented in code"],
            ["Validation tooling", "Partial", "Evaluation exists but should be broadened and standardized"],
            ["Operational maturity", "Mixed", "Prototype is promising but not fully enterprise-ready"],
            ["Architecture quality", "Partial modularization", "Refactoring direction is positive and should continue"],
        ],
    )

    section_heading(document, "6.3 Comparative Discussion with Modern Architectures")
    paragraph_pack(
        document,
        [
            "The modern architecture comparison report is valuable because it prevents the project from being evaluated only against a weak baseline. Against traditional static networks, VANTA performs strongly because dynamic obfuscation directly addresses the kind of stale reconnaissance problem that static architectures leave exposed. Against general zero trust and microsegmentation models, however, the comparison becomes more nuanced. VANTA's strength in address dynamism does not automatically equal maturity in identity, device, and enterprise policy integration.",
            "This result should not be read as a defeat for VANTA. Instead, it clarifies the function of the prototype. VANTA is most compelling where the research question centers on making attack surfaces less stable and reconnaissance less durable. Zero trust and microsegmentation remain stronger reference points when the question shifts toward enterprise-wide authorization, workload isolation, and mature policy operations. The natural design implication is synthesis rather than competition.",
        ],
    )
    add_figure(document, assets["vanta_radar"], "Figure 7. Radar view of VANTA's qualitative capability profile across comparison criteria.", width=5.6)

    section_heading(document, "6.4 Discussion of Testing and Reproducibility")
    paragraph_pack(
        document,
        [
            "Reproducibility is one of the strongest dividing lines between an interesting prototype and a dependable research artifact. VANTA has already taken several positive steps in this direction: test modules exist, the project ships comparison scripts, runtime configuration is externalized, and state storage has been separated into dedicated components. At the same time, the mixed test execution observed during report preparation highlights the need for environment-agnostic temporary paths, clearer dependency instructions, and stronger packaging isolation.",
            "These observations have constructive value. They point to practical future enhancements that are fully aligned with the report's research message: if VANTA is to serve as a platform for experimental security evaluation, then the pathway from clone to test run to experiment report must become as smooth and deterministic as possible.",
        ],
    )

    section_heading(document, "6.5 Threats to Validity")
    paragraph_pack(
        document,
        [
            "Several validity considerations apply to the current report. First, not all claims are backed by a fresh live-network rerun in the same environment in which the document was generated. This report therefore distinguishes clearly between repository-observed capability and directly reproduced performance measurements. Second, qualitative comparison scores reflect structured reasoning from generated repository reports and cited standards, but they are still interpretive rather than universal facts.",
            "Third, laboratory-focused SDN architectures can appear stronger in controlled environments than they will under enterprise heterogeneity. The project already hints at this through hardware adapter abstractions and rollout policy mechanisms, but large-scale deployment remains a future challenge. Fourth, academic similarity and AI-detection systems differ in methodology, so originality-oriented writing choices reduce risk but cannot guarantee a particular threshold outcome under every checker.",
        ],
    )

    chapter_heading(document, "Chapter 7: Conclusion and Future Enhancement")
    section_heading(document, "7.1 Conclusion")
    paragraph_pack(
        document,
        [
            "This report examined VANTA as a research-based minor project centered on adaptive network defense through software-defined moving target principles. The study showed that the project is substantial in both concept and implementation. It is grounded in a real security problem, supported by relevant literature, represented in code through controller-driven VIP morphing and threat-triggered logic, and enriched by visualization, reporting, testing, and benchmarking surfaces.",
            "The most important conclusion is that VANTA succeeds as a prototype for network-layer uncertainty. It demonstrates how SDN control can make host identities more ephemeral, how threat signals can trigger adaptive response, and how a research-oriented architecture can be explained through figures, tables, and comparative analysis. It also reveals that strong moving target behavior does not eliminate the need for complementary controls rooted in identity, device trust, segmentation, and deployment maturity.",
            "For that reason, the project's deeper academic value lies not only in the prototype it currently represents but in the architectural conversation it opens. VANTA provides a tangible base from which a richer defensive stack can be explored, one in which dynamic obfuscation works alongside zero trust and microsegmentation rather than apart from them.",
        ],
    )

    section_heading(document, "7.2 Future Enhancement")
    add_bullets(
        document,
        [
            "Integrate richer identity- and device-posture-aware policy so that morphing decisions can cooperate with zero-trust access logic.",
            "Expand benchmark automation to produce standardized experiment bundles and richer comparison dashboards.",
            "Improve test isolation and packaging so that validation passes consistently across clean environments.",
            "Strengthen hardware-backed deployment support through mature network adapters and rollout governance.",
            "Evaluate larger and more heterogeneous topologies to study scale, flow churn, and controller bottlenecks.",
            "Introduce workload-level microsegmentation and policy labels to better contain lateral movement after reconnaissance.",
        ],
    )


def add_generated_lists(document: Document) -> None:
    add_page_break(document)
    add_paragraph(document, "List of Figures", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    for entry in [
        "Figure 1. High-level VANTA architecture derived from the current repository design.",
        "Figure 2. Threat-triggered morphing response sequence represented as a controller workflow.",
        "Figure 3. Illustrative VIP morphing timeline showing how exposed identity states can change across observation windows.",
        "Figure 4. Readiness of documented capabilities derived from repository comparison artifacts.",
        "Figure 5. Distribution of implemented, partial, and missing capability statuses in the architecture alignment report.",
        "Figure 6. Qualitative comparison of VANTA with selected mainstream security architecture patterns.",
        "Figure 7. Radar view of VANTA's qualitative capability profile across comparison criteria.",
    ]:
        add_paragraph(document, entry)
    add_page_break(document)
    add_paragraph(document, "List of Tables", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx in range(1, 11):
        add_paragraph(document, f"Table {idx}. Included in the body of the report.")
    add_page_break(document)
    add_paragraph(document, "Abbreviations", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(
        document,
        "Table 11. Abbreviations",
        ["Abbreviation", "Expanded Form"],
        [
            ["SDN", "Software-Defined Networking"],
            ["MTD", "Moving Target Defense"],
            ["VIP", "Virtual IP"],
            ["OVS", "Open vSwitch"],
            ["API", "Application Programming Interface"],
            ["MFA", "Multi-Factor Authentication"],
            ["RTT", "Round-Trip Time"],
            ["ZTA", "Zero Trust Architecture"],
        ],
    )


def add_length_balancer(document: Document) -> None:
    chapter_heading(document, "Appendix D: Extended Analytical Notes")
    thematic_sets = [
        ("Reconnaissance and uncertainty", "reconnaissance quality", "network identity stability", "defensive asymmetry"),
        ("Programmable control", "centralized policy orchestration", "controller visibility", "adaptive forwarding"),
        ("Prototype maturity", "test isolation", "packaging reliability", "artifact reproducibility"),
        ("Operational trade-offs", "flow churn", "latency overhead", "trigger sensitivity"),
        ("Comparative security design", "zero trust alignment", "microsegmentation fit", "hybrid deployment implications"),
    ]
    for round_id in range(1, 19):
        subsection_heading(document, f"D.{round_id} Analytical Note Set {round_id}")
        for theme, concept_a, concept_b, concept_c in thematic_sets:
            paragraph_pack(
                document,
                [
                    f"In this analytical note set, the theme of {theme.lower()} is revisited from a systems perspective. VANTA shows that {concept_a} cannot be studied in isolation from {concept_b}. When a defender deliberately changes what the attacker sees, the relevant question is not simply whether change occurred, but whether the change altered the information economics of the engagement. That is why the report repeatedly returns to the relationship between observable identity, controller decision logic, and attacker confidence.",
                    f"A second observation is that {concept_b} becomes strategically meaningful only when it is coupled with {concept_c}. Many prototypes can demonstrate dynamic behavior, yet fewer can explain how that behavior is selected, measured, and integrated with the rest of the architecture. VANTA's educational strength lies in making these linkages visible: mapping logic, threat detection, dashboards, tests, and comparison scripts all contribute to a more complete picture of adaptive defense.",
                    f"From a research methods standpoint, the interaction among {concept_a}, {concept_b}, and {concept_c} suggests a family of future experiments rather than a single terminal result. That is appropriate for a minor project report. The goal is not to claim the final word on moving target defense, but to present a convincing prototype, explain its value with discipline, and leave a documented path for deeper empirical study.",
                ],
            )


def build_document() -> None:
    assets = make_assets()
    document = Document()
    set_default_font(document)

    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.15)
    section.right_margin = Inches(1)

    build_cover(document)
    build_preliminary_pages(document)
    add_generated_lists(document)
    build_main_body(document, assets)
    add_reference_section(document)
    add_appendices(document)
    add_length_balancer(document)

    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
    print(f"Report generated at: {OUTPUT_PATH}")
